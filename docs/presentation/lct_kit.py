# -*- coding: utf-8 -*-
"""Примитивы оформления по шаблону ЛЦТ 2026 (docs/presentation/template/).

Шаблон задаёт фоны (фиолетовый градиент, светлый, светлый с силуэтом города,
титульный — логотипы организаторов уже на них), шрифт Montserrat и палитру
темы. Здесь — то, чего в шаблоне нет готовым: плашка-заголовок, карточки,
таблицы, списки с розовыми маркерами, снимки DXA, всё в цветах темы.
"""
from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
X0 = Emu(346075)                     # левый край плашки-заголовка в шаблоне
XR = W - Inches(0.4)
CW = XR - X0
I = Inches
FONT = 'Montserrat'


def rgb(h: str) -> RGBColor:
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# Палитра темы шаблона (accent1..6, dk2) и производные для текста
C = {k: rgb(v) for k, v in dict(
    pink='FF0053', pink_soft='FFD6E3', pink2='FC3777', lav='8A83D1', purple='520977', deep='2D1451',
    ink='1C1D22', ink2='4A4458', ink3='8B8599', white='FFFFFF', lilac='DCCEE4', line='C9C3DA',
    film='120A1C', film_ink='D9D2E6', geo='8A83D1', ok='1F8A5B').items()}


# --------------------------------------------------------------------------- #
#  Текст
# --------------------------------------------------------------------------- #
def _run(p, text, size, color, bold=False, italic=False, font=FONT):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', font)
    return r


def text(slide, x, y, w, h, content, size=14, color=None, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1, after=0, caps=False, italic=False, tracking=0):
    """content: строка, список абзацев или список абзацев из кусков (текст, {опции})."""
    color = C['ink'] if color is None else color
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', 0)
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        chunks = para if isinstance(para, list) else [(para, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            t = t.upper() if o.get('caps', caps) else t
            r = _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', bold),
                     o.get('italic', italic))
            sp = o.get('tracking', tracking)
            if sp:
                r._r.get_or_add_rPr().set('spc', str(int(sp * 100)))
    return tb


def bullets(slide, x, y, w, h, items, size=13, color=None, gap=5, marker='▪', marker_color=None,
            spacing=1.12):
    """Список с квадратными розовыми маркерами, как на слайдах шаблона."""
    color = C['ink'] if color is None else color
    tb = text(slide, x, y, w, h, [''], size=size)
    tf = tb.text_frame
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.24)))
        pPr.set('indent', str(-Inches(0.24)))
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        etree.SubElement(buClr, qn('a:srgbClr')).set('val', str(marker_color or C['pink']))
        etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
        etree.SubElement(pPr, qn('a:buChar')).set('char', marker)
        chunks = item if isinstance(item, list) else [(item, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', False))
    return tb


# --------------------------------------------------------------------------- #
#  Фигуры
# --------------------------------------------------------------------------- #
def _no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn('a:effectLst')) is None:
        etree.SubElement(spPr, qn('a:effectLst'))


def rect(slide, x, y, w, h, fill=None, line=None, radius=0.0, line_w=0.75, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   x, y, w, h)
    if radius:
        shape.adjustments[0] = radius
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        if alpha is not None:
            clr = shape.fill._xPr.find(qn('a:solidFill'))[0]
            etree.SubElement(clr, qn('a:alpha')).set('val', str(int(alpha * 100000)))
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    _no_shadow(shape)
    shape.text_frame.text = ''
    return shape


def oval(slide, cx, cy, d, fill, line=None, line_w=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, cy - d // 2, d, d)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    _no_shadow(s)
    return s


def line(slide, x1, y1, x2, y2, color, width=1.0, dash=None, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn('a:prstDash')).set('val', dash)
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle')
        te.set('w', 'med')
        te.set('len', 'med')
    return c


def label(slide, x, y, w, text_, color=None, size=10):
    """Надпись-«рубрика» над блоком: капс, разрядка, розовая."""
    return text(slide, x, y, w, I(0.25), text_, size=size, bold=True, color=color or C['pink'],
                caps=True, tracking=0.8)


def card(slide, x, y, w, h, title='', body='', kicker='', dark=True, title_size=16, body_size=12.5,
         metric='', accent=False, fill=None):
    """Белая карточка. На тёмном фоне — без рамки (как в шаблоне), на светлом —
    с тонкой сиреневой рамкой."""
    border = C['pink'] if accent else (None if dark else C['line'])
    rect(slide, x, y, w, h, fill=fill or C['white'], line=border, radius=0.06 if h < I(2.2) else 0.045,
         line_w=1.5 if accent else 0.75)
    pad = I(0.24)
    cy = y + I(0.2)
    if kicker:
        label(slide, x + pad, cy, w - 2 * pad, kicker, size=9.5)
        cy += I(0.3)
    if title:
        lines = title_lines(title, (w - 2 * pad) / 914400, title_size)
        text(slide, x + pad, cy, w - 2 * pad, Pt(title_size * 1.2 * lines), title, size=title_size, bold=True,
             color=C['purple'], spacing=1.0)
        cy += I(0.1) + Pt(title_size * 1.2 * lines)
    if body:
        items = body if isinstance(body, list) else None
        bh = y + h - cy - (I(0.55) if metric else I(0.15))
        if items:
            bullets(slide, x + pad, cy, w - 2 * pad, bh, items, size=body_size, color=C['ink2'])
        else:
            text(slide, x + pad, cy, w - 2 * pad, bh, body, size=body_size, color=C['ink2'], spacing=1.12)
    if metric:
        line(slide, x + pad, y + h - I(0.5), x + w - pad, y + h - I(0.5), C['pink_soft'], 1)
        text(slide, x + pad, y + h - I(0.42), w - 2 * pad, I(0.3), metric, size=11.5, bold=True,
             color=C['pink'])


def title_lines(title: str, width_in: float, size: float) -> int:
    """Сколько строк займёт жирный заголовок Montserrat: средняя ширина знака
    около 0.62 кегля — оценка с запасом, чтобы текст под ним не наезжал."""
    import math
    per_line = max(1, int(width_in * 72 / (size * 0.62)))
    n = 0
    for part in title.split('\n'):
        words, cur, k = part.split(), 0, 1
        for wd in words:
            if cur and cur + 1 + len(wd) > per_line:
                k += 1
                cur = len(wd)
            else:
                cur += (1 if cur else 0) + len(wd)
        n += k
    return max(1, n) if math.isfinite(per_line) else 1


def stat(slide, x, y, w, value, caption, color=None, cap_color=None, size=34, unit=''):
    text(slide, x, y, w, I(0.62), [[(value, {'size': size, 'bold': True}),
                                    (unit, {'size': size * 0.45, 'bold': True})]],
         color=color or C['pink'])
    text(slide, x, y + Pt(size * 1.2), w, I(0.6), caption, size=11.5,
         color=cap_color or C['ink2'], spacing=1.05)


def number_badge(slide, x, y, n, color=None, size=I(0.46)):
    """Квадрат с цифрой — как на слайде «Стадии» шаблона."""
    b = rect(slide, x, y, size, size, fill=C['white'], line=color or C['pink'], radius=0.12, line_w=1.25)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, str(n), 15, color or C['pink'], bold=True)
    return b


# --------------------------------------------------------------------------- #
#  Изображения
# --------------------------------------------------------------------------- #
def picture_fit(slide, path: Path, x, y, w, h, frame=True, caption=None, cap_h=I(0.55), cap_size=10.5,
                cap_color=None, radius=0.04):
    """Снимок DXA на тёмной подложке, вписан с сохранением пропорций."""
    if frame:
        rect(slide, x, y, w, h, fill=C['film'], radius=radius)
    ih = h - (cap_h if caption else 0)
    with Image.open(path) as im:
        iw_px, ih_px = im.size
    k = min(w / iw_px, ih / ih_px)
    pw, ph = int(iw_px * k), int(ih_px * k)
    pic = slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (ih - ph) // 2, pw, ph)
    if caption:
        text(slide, x + I(0.16), y + ih + I(0.07), w - I(0.32), cap_h - I(0.1), caption, size=cap_size,
             color=cap_color or C['film_ink'], spacing=1.08)
    return pic


def picture(slide, path: Path, x, y, w=None, h=None):
    with Image.open(path) as im:
        iw, ih = im.size
    if w is not None and h is None:
        h = int(w * ih / iw)
    elif h is not None and w is None:
        w = int(h * iw / ih)
    return slide.shapes.add_picture(str(path), x, y, w, h)


# --------------------------------------------------------------------------- #
#  Таблица
# --------------------------------------------------------------------------- #
def _cell_border(cell, top=None, bottom=None):
    """lnL/lnR/lnT/lnB идут в tcPr первыми — иначе PowerPoint их игнорирует."""
    tcPr = cell._tc.get_or_add_tcPr()
    for i, (tag, spec) in enumerate((('a:lnL', None), ('a:lnR', None), ('a:lnT', top), ('a:lnB', bottom))):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        tcPr.insert(i, ln)
        if spec is None:
            ln.set('w', '0')
            etree.SubElement(ln, qn('a:noFill'))
        else:
            color, width = spec
            ln.set('w', str(int(Pt(width))))
            sf = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr')).set('val', str(color))


def table(slide, x, y, w, rows, widths, size=12, header_size=10, row_h=I(0.38), align=None,
          highlight=(), bold_last=False, color=None):
    """rows[0] — заголовок. widths — доли. align — 'l'/'r'/'c' на колонку."""
    n, m = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n, m, x, y, w, row_h * n)
    tbl = shape.table
    style = tbl._tbl.tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'       # без стиля и сетки
    tbl.first_row, tbl.horz_banding = False, False
    total = sum(widths)
    for j, f in enumerate(widths):
        tbl.columns[j].width = int(w * f / total)
    align = align or ['l'] + ['r'] * (m - 1)
    amap = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}
    ink = color or C['ink']
    for i, row in enumerate(rows):
        tbl.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.background()
            if i in highlight:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C['pink_soft']
            cell.margin_left = I(0.07) if j else I(0.05)
            cell.margin_right = I(0.08)
            cell.margin_top = cell.margin_bottom = I(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = amap[align[j]]
            if i == 0:
                _run(p, str(val).upper(), header_size, C['purple'], bold=True)
                _cell_border(cell, bottom=(C['purple'], 1.5))
            else:
                chunks = val if isinstance(val, list) else [(str(val), {})]
                for chunk in chunks:
                    t, o = (chunk, {}) if isinstance(chunk, str) else chunk
                    _run(p, t, o.get('size', size), o.get('color', ink),
                         o.get('bold', bold_last and i == n - 1))
                _cell_border(cell, bottom=(C['line'], 0.75))
    return shape


# --------------------------------------------------------------------------- #
#  Колода по шаблону
# --------------------------------------------------------------------------- #
STYLE_SRC = {'light': 1, 'dark': 2, 'city': 3}      # индексы слайдов шаблона с нужными фонами


class TemplateDeck:
    """Колода на основе шаблона: слайды шаблона 0–3 (титул, о команде, команда,
    история) заполняются как есть, остальные добавляются с фонами шаблона."""

    def __init__(self, template: Path):
        self.prs = Presentation(str(template))
        self.src = list(self.prs.slides)
        self.layout = next(l for l in self.prs.slide_layouts if l.name == 'Пустой с заголовком')
        self.title_layout = next(l for l in self.prs.slide_layouts if l.name == 'Титульный слайд')
        self._bg = {}
        for style, idx in STYLE_SRC.items():
            s = self.src[idx]
            bg = s._element.find(qn('p:cSld')).find(qn('p:bg'))
            blip = bg.find('.//' + qn('a:blip'))
            self._bg[style] = (bg, s.part.related_part(blip.get(qn('r:embed'))))

    # ------------------------------------------------------------------ фон
    def set_bg(self, slide, style: str):
        bg_el, img_part = self._bg[style]
        new = copy.deepcopy(bg_el)
        rid = slide.part.relate_to(img_part, RT.IMAGE)
        new.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
        cSld = slide._element.find(qn('p:cSld'))
        old = cSld.find(qn('p:bg'))
        if old is not None:
            cSld.remove(old)
        cSld.insert(0, new)

    def slide_number(self, slide, style: str):
        color = 'FFFFFF' if style in ('dark', 'title') else '520977'
        xml = (f'<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
               f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
               f'<p:nvSpPr><p:cNvPr id="{900 + len(slide.shapes)}" name="Номер слайда"/>'
               f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="sldNum" sz="quarter" idx="4"/>'
               f'</p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="{int(I(12.1))}" y="{int(I(6.95))}"/>'
               f'<a:ext cx="{int(I(0.95))}" cy="{int(I(0.4))}"/></a:xfrm></p:spPr>'
               f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr algn="r"/>'
               f'<a:fld id="{{52DCC5B9-D646-4B76-891F-6FF4E74E9CB4}}" type="slidenum">'
               f'<a:rPr lang="ru-RU" sz="1200"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
               f'<a:latin typeface="{FONT}"/></a:rPr><a:t>‹#›</a:t></a:fld></a:p></p:txBody></p:sp>')
        slide.shapes._spTree.append(etree.fromstring(xml))

    # ------------------------------------------------------------------ слайды
    def new(self, style: str, section: str, headline: str = '', pill=None):
        """Слайд с фоном шаблона, плашкой-разделом и заголовком-утверждением."""
        s = self.prs.slides.add_slide(self.layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        self.set_bg(s, style)
        dark = style == 'dark'
        fill = pill or (C['purple'] if style == 'light' else C['pink'])
        self.pill(s, section, fill)
        if headline:
            text(s, X0, I(1.2), I(12.4), I(0.9), headline, size=22, bold=True, spacing=1.02,
                 color=C['white'] if dark else C['purple'])
        self.slide_number(s, style)
        return s

    def closing(self):
        s = self.prs.slides.add_slide(self.title_layout)
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        self.slide_number(s, 'title')
        return s

    @staticmethod
    def pill(slide, section: str, fill):
        sec = section.upper()
        w = min(I(0.75 + 0.205 * len(sec)), I(7.0))
        p = rect(slide, X0, Emu(324562), w, Emu(620919), fill=fill, radius=0.187)
        tf = p.text_frame
        tf.margin_left, tf.margin_right = I(0.3), I(0.2)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        _run(para, sec, 20, C['white'], bold=True)
        return p

    def notes(self, slide, paragraphs: list[str]):
        slide.notes_slide.notes_text_frame.text = '\n\n'.join(paragraphs)

    def save(self, path: Path):
        self.prs.save(str(path))
