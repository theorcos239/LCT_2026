# -*- coding: utf-8 -*-
"""Примитивы оформления для нативной PPTX: цвета, шрифты, текст, карточки, таблицы.

Тот же визуальный язык, что у HTML-версии: светлый «бумажный» фон, тёмные
плёночные панели под снимки DXA, один акцентный синий, моноширинные цифры.
Шрифты — из комплекта Windows/Office (Segoe UI, Consolas), чтобы презентация
выглядела одинаково на любом ноутбуке жюри без установки шрифтов.
"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.62)                    # поля слева и справа
CW = W - 2 * MX                      # ширина контента


def rgb(h: str) -> RGBColor:
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


C = {k: rgb(v) for k, v in dict(
    paper='F2F4F5', sheet='FFFFFF', ink='111C24', ink2='44545F', ink3='7A8995', rule='D3DBE0',
    rule2='E6EBEE', film='0B1015', film_ink='C9D3DA', accent='1B5FD0', accent_soft='E4ECFA',
    accent_ink='123F8C', geo='8C9AA5', ok='1F8A5B', ok_soft='E1F2EA', bad='C9443B',
    bad_soft='F8E3E1', warn='A8790A', heat='FF6000', cnn='6F9BE6').items()}

SANS, SANS_B, MONO = 'Segoe UI', 'Segoe UI Semibold', 'Consolas'


# --------------------------------------------------------------------------- #
#  Текст
# --------------------------------------------------------------------------- #
def _run(p, text, size, color, bold=False, font=SANS, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:cs', 'a:ea'):                         # кириллица и прочие скрипты тем же шрифтом
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', font)
    return r


def text(slide, x, y, w, h, content, size=16, color=None, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, after=0, caps=False,
         tracking=0):
    """content: строка, список абзацев или список абзацев из кусков (текст, {опции})."""
    color = color or C['ink']
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', 0)
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        chunks = para if isinstance(para, list) else [(para, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            t = t.upper() if (o.get('caps', caps)) else t
            r = _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', bold),
                     o.get('font', font), o.get('italic', False))
            sp = o.get('tracking', tracking)
            if sp:
                r._r.get_or_add_rPr().set('spc', str(int(sp * 100)))
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=None, gap=6, marker='–',
            marker_color=None, spacing=1.2):
    color = color or C['ink']
    tb = text(slide, x, y, w, h, [''], size=size)
    tf = tb.text_frame
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.26)))
        pPr.set('indent', str(-Inches(0.26)))
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        clr = etree.SubElement(buClr, qn('a:srgbClr'))
        clr.set('val', str(marker_color or C['accent']))
        buFont = etree.SubElement(pPr, qn('a:buFont'))
        buFont.set('typeface', SANS)
        bu = etree.SubElement(pPr, qn('a:buChar'))
        bu.set('char', marker)
        chunks = item if isinstance(item, list) else [(item, {})]
        for chunk in chunks:
            t, o = (chunk, {}) if isinstance(chunk, str) else chunk
            _run(p, t, o.get('size', size), o.get('color', color), o.get('bold', False),
                 o.get('font', SANS))
    return tb


# --------------------------------------------------------------------------- #
#  Фигуры
# --------------------------------------------------------------------------- #
def _no_shadow(shape):
    spPr = shape._element.spPr
    if spPr.find(qn('a:effectLst')) is None:
        etree.SubElement(spPr, qn('a:effectLst'))


def rect(slide, x, y, w, h, fill=None, line=None, radius=0.0, line_w=0.75):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   x, y, w, h)
    if radius:
        shape.adjustments[0] = radius
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    _no_shadow(shape)
    shape.text_frame.text = ''
    return shape


def oval(slide, cx, cy, d, fill, line=None, line_w=1.5):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, cy - d // 2, d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    _no_shadow(s)
    return s


def line(slide, x1, y1, x2, y2, color, width=1.0, dash=None, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', dash)
    if arrow:
        te = etree.SubElement(ln, qn('a:tailEnd'))
        te.set('type', 'triangle')
        te.set('w', 'med')
        te.set('len', 'med')
    return c


def card(slide, x, y, w, h, kicker='', title='', body='', accent=False, body_size=14,
         metric='', fill=None, title_size=17):
    rect(slide, x, y, w, h, fill=fill or C['sheet'], line=C['accent'] if accent else C['rule'],
         radius=0.05, line_w=1.25 if accent else 0.75)
    pad = Inches(0.22)
    cy = y + Inches(0.18)
    if kicker:
        text(slide, x + pad, cy, w - 2 * pad, Inches(0.25), kicker, size=10, font=MONO,
             color=C['accent'] if accent else C['ink3'], caps=True, tracking=0.6)
        cy += Inches(0.3)
    if title:
        text(slide, x + pad, cy, w - 2 * pad, Inches(0.4), title, size=title_size, font=SANS_B)
        cy += Inches(0.42)
    if body:
        text(slide, x + pad, cy, w - 2 * pad, y + h - cy - Inches(0.5 if metric else 0.15), body,
             size=body_size, color=C['ink2'], spacing=1.2)
    if metric:
        line(slide, x + pad, y + h - Inches(0.52), x + w - pad, y + h - Inches(0.52), C['rule2'], 0.75)
        text(slide, x + pad, y + h - Inches(0.43), w - 2 * pad, Inches(0.3), metric, size=12.5,
             font=MONO, color=C['accent_ink'])


def stat(slide, x, y, w, value, unit, caption, value_size=38, divider=True):
    if divider:
        line(slide, x - Inches(0.18), y, x - Inches(0.18), y + Inches(1.15), C['rule'], 0.75)
    text(slide, x, y, w, Inches(0.7), [[(value, {'font': MONO, 'size': value_size}),
                                        (unit, {'font': MONO, 'size': value_size * 0.45,
                                                'color': C['ink2']})]])
    text(slide, x, y + Inches(0.72), w, Inches(0.55), caption, size=12, color=C['ink2'], spacing=1.1)


def pill(slide, x, y, label, fg, bg, size=10.5):
    w = Inches(0.12 + 0.075 * len(label))
    s = rect(slide, x, y, w, Inches(0.26), fill=bg, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, label, size, fg, font=MONO)
    return w


# --------------------------------------------------------------------------- #
#  Изображения
# --------------------------------------------------------------------------- #
def picture_fit(slide, path: Path, x, y, w, h, film=True, caption=None, cap_h=Inches(0.5),
                cap_size=10.5):
    """Снимок на тёмной «плёнке», вписан с сохранением пропорций."""
    if film:
        rect(slide, x, y, w, h, fill=C['film'], radius=0.03)
    ih = h - (cap_h if caption else 0)
    with Image.open(path) as im:
        iw_px, ih_px = im.size
    k = min(w / iw_px, ih / ih_px)
    pw, ph = int(iw_px * k), int(ih_px * k)
    pic = slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (ih - ph) // 2, pw, ph)
    if caption:
        line(slide, x, y + ih, x + w, y + ih, rgb('1F2A33'), 0.75)
        text(slide, x + Inches(0.14), y + ih + Inches(0.08), w - Inches(0.28), cap_h - Inches(0.1),
             caption, size=cap_size, font=MONO, color=C['film_ink'], spacing=1.15)
    return pic


# --------------------------------------------------------------------------- #
#  Таблица
# --------------------------------------------------------------------------- #
def _cell_border(cell, top=None, bottom=None):
    """Границы ячейки. В схеме OOXML lnL/lnR/lnT/lnB идут ПЕРВЫМИ в tcPr, до заливки —
    иначе PowerPoint молча их игнорирует."""
    tcPr = cell._tc.get_or_add_tcPr()
    for i, (tag, spec) in enumerate((('a:lnL', None), ('a:lnR', None), ('a:lnT', top), ('a:lnB', bottom))):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        tcPr.insert(i, ln)
        if spec is None:
            ln.set('w', '0')
            etree.SubElement(ln, qn('a:noFill'))
        else:
            color, width = spec
            ln.set('w', str(int(Pt(width))))
            sf = etree.SubElement(ln, qn('a:solidFill'))
            c = etree.SubElement(sf, qn('a:srgbClr'))
            c.set('val', str(color))


def table(slide, x, y, w, rows, widths, size=12, header_size=9.5, row_h=Inches(0.36),
          align=None, highlight=(), mono_cols=(), bold_last=False):
    """rows[0] — заголовок. widths — доли. align — 'l'/'r' на колонку."""
    n, m = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n, m, x, y, w, row_h * n)
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn('a:tableStyleId'))
    if style is not None:
        style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'     # «без стиля, без сетки»
    tbl.first_row, tbl.horz_banding = False, False
    total = sum(widths)
    for j, f in enumerate(widths):
        tbl.columns[j].width = int(w * f / total)
    align = align or ['l'] + ['r'] * (m - 1)
    for i, row in enumerate(rows):
        tbl.rows[i].height = row_h
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.background()
            if i in highlight:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C['accent_soft']
            cell.margin_left = Inches(0.06) if j else Inches(0.04)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if align[j] == 'r' else PP_ALIGN.LEFT
            if i == 0:
                _run(p, str(val).upper(), header_size, C['ink3'], font=MONO)
                _cell_border(cell, bottom=(C['ink'], 1.25))
            else:
                chunks = val if isinstance(val, list) else [(str(val), {})]
                for chunk in chunks:
                    t, o = (chunk, {}) if isinstance(chunk, str) else chunk
                    _run(p, t, o.get('size', size), o.get('color', C['ink']),
                         o.get('bold', bold_last and i == n - 1),
                         o.get('font', MONO if j in mono_cols else SANS))
                _cell_border(cell, bottom=(C['rule'], 0.75))
    return shape


# --------------------------------------------------------------------------- #
#  Каркас слайда
# --------------------------------------------------------------------------- #
def frame(prs, n: int, total: int, section: str = '', title: str = '', dark_right=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = C['paper']
    if section:
        text(slide, MX, Inches(0.42), Inches(8), Inches(0.3),
             [[(f'{n:02d}   ', {'color': C['ink3']}), (section, {'color': C['accent']})]],
             size=11, font=MONO, caps=True, tracking=1.0)
    if title:
        text(slide, MX, Inches(0.72), Inches(9.6), Inches(1.0), title, size=28, font=SANS_B,
             spacing=1.02)
    # подвал и «миллиметровая линейка» по нижнему краю
    text(slide, MX, H - Inches(0.42), Inches(6), Inches(0.25), 'Контроль качества DXA · ЛЦТ 2026',
         size=9.5, font=MONO, color=C['ink3'])
    text(slide, W - MX - Inches(1.5), H - Inches(0.42), Inches(1.5), Inches(0.25),
         f'{n:02d} / {total:02d}', size=9.5, font=MONO, color=C['ink3'], align=PP_ALIGN.RIGHT)
    slide.shapes.add_picture(str(ruler_png()), 0, H - Inches(0.12), W, Inches(0.12))
    return slide


def ruler_png() -> Path:
    """Миллиметровая линейка по нижнему краю — одной картинкой, а не сотней фигур."""
    from PIL import ImageDraw
    p = Path(__file__).resolve().parent / 'assets' / 'ruler.png'
    if p.exists():
        return p
    dpi = 200
    wpx, hpx = int(13.333 * dpi), int(0.12 * dpi)
    im = Image.new('RGBA', (wpx, hpx), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    step = 0.1 * dpi
    i = 0
    while i * step < wpx:
        x = int(i * step)
        long_ = i % 5 == 0
        col = (122, 137, 149, 255) if long_ else (211, 219, 224, 255)
        d.line([(x, hpx - (hpx if long_ else hpx // 2)), (x, hpx)], fill=col, width=2)
        i += 1
    p.parent.mkdir(exist_ok=True)
    im.save(p)
    return p


def notes(slide, paragraphs: list[str]) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = '\n\n'.join(paragraphs)
