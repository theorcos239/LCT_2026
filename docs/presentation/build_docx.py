# -*- coding: utf-8 -*-
"""Доклад к презентации по слайдам: DXA_QC_speech.docx.

    python docs/presentation/build_docx.py

Текст — speech.py (тот же, что в заметках докладчика PPTX), миниатюры
слайдов — PNG, которые экспортирует PowerPoint (build_pptx.py).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1]))

import build  # noqa: E402
import speech  # noqa: E402

INK = RGBColor(0x11, 0x1C, 0x24)
INK2 = RGBColor(0x44, 0x54, 0x5F)
INK3 = RGBColor(0x7A, 0x89, 0x95)
ACCENT = RGBColor(0x1B, 0x5F, 0xD0)
FONT, MONO = 'Segoe UI', 'Consolas'
PNG = HERE / 'assets' / 'pptx_png'


def _font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    return run


def para(doc, text='', size=11, color=INK, bold=False, name=FONT, after=6, before=0, align=None,
         line=1.25, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), line
    if align:
        p.alignment = align
    if text:
        _font(p.add_run(text), name, size, color, bold, italic)
    return p


def rule(doc, color='D3DBE0', size=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def seconds(t: str) -> int:
    # «0:15 + 3:00» — слайд и демонстрация: в регламент речи идёт только первое
    m = re.search(r'(\d+):(\d+)', t)
    return int(m.group(1)) * 60 + int(m.group(2)) if m else 0


def main() -> int:
    D = build.load()
    cases = build.render_cases()
    S = speech.slides(D, cases)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles['Normal']
    st.font.name, st.font.size = FONT, Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    # колонтитул
    fp = sec.footer.paragraphs[0]
    _font(fp.add_run('Контроль качества DXA · доклад к презентации · ЛЦТ 2026'), MONO, 8, INK3)

    # титул
    para(doc, 'ЛЦТ 2026 · КЕЙС ДЕПАРТАМЕНТА ЗДРАВООХРАНЕНИЯ МОСКВЫ', 9, ACCENT, name=MONO, before=60, after=10)
    para(doc, 'Контроль качества DXA', 30, INK, bold=True, after=4, line=1.0)
    para(doc, 'Доклад к презентации по слайдам', 16, INK2, after=18)
    para(doc, 'Сервис проверяет укладку и разметку денситометрии до того, как исследование уйдёт к врачу, '
              'и объясняет каждое замечание в миллиметрах и градусах.', 12, INK2, after=18, line=1.35)
    total = sum(seconds(x['time']) for x in S)
    rule(doc)
    para(doc, f'Регламент: около {round(total / 60)} минут выступления и 3 минуты демонстрации. Для каждого слайда '
              'ниже — текст для чтения вслух и короткий список «если спросят»: факты, которые в регламент '
              'не помещаются, но нужны в ответах на вопросы. Все числа взяты из файлов метрик репозитория '
              'и совпадают со слайдами и README.', 10.5, INK2, after=12, line=1.35)

    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for c, txt in zip(hdr, ('№', 'слайд', 'время')):
        _font(c.paragraphs[0].add_run(txt.upper()), MONO, 8.5, INK3)
    for i, sl in enumerate(S, 1):
        r = t.add_row().cells
        _font(r[0].paragraphs[0].add_run(f'{i:02d}'), MONO, 10, INK3)
        _font(r[1].paragraphs[0].add_run(sl['title']), FONT, 10.5, INK)
        _font(r[2].paragraphs[0].add_run(sl['time']), MONO, 10, INK2)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Cm(1.2), Cm(12.5), Cm(3.2)

    # слайды
    for i, sl in enumerate(S, 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = para(doc, after=2)
        _font(p.add_run(f'СЛАЙД {i:02d}   '), MONO, 9, INK3)
        _font(p.add_run(sl['time']), MONO, 9, ACCENT)
        para(doc, sl['title'], 18, INK, bold=True, after=8, line=1.05)
        img = PNG / f'{i:02d}.png'
        if img.exists():
            pic = doc.add_paragraph()
            pic.paragraph_format.space_after = Pt(10)
            pic.add_run().add_picture(str(img), width=Cm(17))
        para(doc, 'ТЕКСТ ВЫСТУПЛЕНИЯ', 8.5, ACCENT, name=MONO, after=4)
        for t_ in sl['text']:
            para(doc, t_, 11.5, INK, after=8, line=1.4)
        if sl.get('ask'):
            para(doc, 'ЕСЛИ СПРОСЯТ', 8.5, INK3, name=MONO, before=6, after=4)
            for a in sl['ask']:
                bp = para(doc, after=4, line=1.3)
                bp.paragraph_format.left_indent = Cm(0.5)
                bp.paragraph_format.first_line_indent = Cm(-0.5)
                _font(bp.add_run('–  '), FONT, 10.5, ACCENT)
                _font(bp.add_run(a), FONT, 10.5, INK2)

    # вопросы
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, 'ПОСЛЕ ДЕМОНСТРАЦИИ', 9, ACCENT, name=MONO, after=2)
    para(doc, 'Ожидаемые вопросы и ответы', 18, INK, bold=True, after=12)
    for q, a in speech.QA:
        para(doc, q, 12, INK, bold=True, after=3, before=6)
        para(doc, a, 11, INK2, after=8, line=1.35)

    out = HERE / 'DXA_QC_speech.docx'
    doc.save(str(out))
    print(out)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
